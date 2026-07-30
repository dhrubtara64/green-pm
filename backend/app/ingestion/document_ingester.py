"""
Document folder ingestion — PDF, DOCX, XLSX, TXT.
Extracts text, classifies content via LLM, and links to Activity/Deliverable entities.
"""
import io
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path

from sqlalchemy.orm import Session

from app.models.evidence_item import EvidenceItem, SourceReliabilitySignal, RelationType
from app.models.associations import activity_evidence_items
from app.agents.evidence_agent import classify_and_link_document


def _extract_text_pdf(content: bytes, filename: str) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception as e:
        return f"[PDF extraction failed: {e}]"


def _extract_text_docx(content: bytes, filename: str) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs)
    except Exception as e:
        return f"[DOCX extraction failed: {e}]"


def _extract_text_xlsx(content: bytes, filename: str) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"=== Sheet: {sheet.title} ===")
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join(str(c) if c is not None else "" for c in row)
                if row_text.strip():
                    lines.append(row_text)
        return "\n".join(lines)
    except Exception as e:
        return f"[XLSX extraction failed: {e}]"


def _extract_text_txt(content: bytes, filename: str) -> str:
    return content.decode("utf-8", errors="replace")


def _extract_text(content: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return _extract_text_pdf(content, filename)
    elif ext in (".docx", ".doc"):
        return _extract_text_docx(content, filename)
    elif ext in (".xlsx", ".xls"):
        return _extract_text_xlsx(content, filename)
    else:
        return _extract_text_txt(content, filename)


def _infer_reliability(filename: str, extracted_text: str) -> SourceReliabilitySignal:
    """Heuristic source reliability signal at ingestion time."""
    name_lower = filename.lower()
    text_lower = extracted_text.lower()

    # Official approved documents
    approval_keywords = ["approved", "certified", "signed", "rev ", "revision", "transmittal"]
    if any(k in name_lower for k in ["approved", "certified", "signed", "afc"]):
        return SourceReliabilitySignal.high
    if any(k in text_lower for k in approval_keywords) and "drawing" in text_lower:
        return SourceReliabilitySignal.high

    # Progress reports and internal trackers
    report_keywords = ["weekly report", "progress report", "status update", "inspection", "testing"]
    if any(k in text_lower for k in report_keywords):
        return SourceReliabilitySignal.medium

    # Informal updates
    if any(k in name_lower for k in ["draft", "wip", "temp", "informal"]):
        return SourceReliabilitySignal.low

    return SourceReliabilitySignal.unverified


def ingest_document_folder(
    folder_path: str,
    db: Session,
    use_llm: bool = True,
) -> dict:
    """
    Ingests all documents in a folder, extracts text, classifies each,
    and links resulting EvidenceItems to Activities/Deliverables.
    """
    folder = Path(folder_path)
    if not folder.exists() or not folder.is_dir():
        return {"error": f"Folder not found: {folder_path}"}

    supported_extensions = {".pdf", ".docx", ".doc", ".xlsx", ".xls", ".txt"}
    items_created = 0
    items_linked = 0
    errors = []

    for filepath in sorted(folder.iterdir()):
        if filepath.suffix.lower() not in supported_extensions:
            continue

        try:
            content = filepath.read_bytes()
            extracted_text = _extract_text(content, filepath.name)
            if not extracted_text.strip():
                continue

            reliability = _infer_reliability(filepath.name, extracted_text)

            # Determine document timestamp from file modification time
            mtime = os.path.getmtime(filepath)
            doc_timestamp = datetime.fromtimestamp(mtime, tz=timezone.utc)

            item_id = f"ev-doc-{uuid.uuid4().hex[:12]}"
            evidence_item = EvidenceItem(
                id=item_id,
                source_system="document_folder",
                ingesting_connector="document_folder_ingestor_v1",
                provenance_ref=str(filepath),
                extracted_content=extracted_text[:4000],  # cap for storage
                source_excerpt=extracted_text[:500],
                relation_type=RelationType.supports_progress_of,
                source_reliability_signal=reliability,
                timestamp=doc_timestamp,
            )
            db.add(evidence_item)
            db.flush()
            items_created += 1

            if use_llm:
                linked = classify_and_link_document(
                    evidence_item_id=item_id,
                    filename=filepath.name,
                    extracted_text=extracted_text,
                    db=db,
                )
                items_linked += linked

        except Exception as e:
            errors.append(f"{filepath.name}: {e}")

    db.commit()
    return {
        "folder": folder_path,
        "items_created": items_created,
        "items_linked": items_linked,
        "errors": errors,
    }


def ingest_single_document(
    content: bytes,
    filename: str,
    db: Session,
    use_llm: bool = True,
) -> dict:
    """Ingest a single uploaded document file."""
    extracted_text = _extract_text(content, filename)
    if not extracted_text.strip():
        return {"error": "No text extracted from document"}

    reliability = _infer_reliability(filename, extracted_text)
    item_id = f"ev-doc-{uuid.uuid4().hex[:12]}"

    evidence_item = EvidenceItem(
        id=item_id,
        source_system="document_upload",
        ingesting_connector="document_upload_ingestor_v1",
        provenance_ref=filename,
        extracted_content=extracted_text[:4000],
        source_excerpt=extracted_text[:500],
        relation_type=RelationType.supports_progress_of,
        source_reliability_signal=reliability,
        timestamp=datetime.now(timezone.utc),
    )
    db.add(evidence_item)
    db.flush()

    linked = 0
    if use_llm:
        linked = classify_and_link_document(
            evidence_item_id=item_id,
            filename=filename,
            extracted_text=extracted_text,
            db=db,
        )

    db.commit()
    return {
        "evidence_item_id": item_id,
        "items_linked": linked,
        "reliability": reliability.value,
    }
